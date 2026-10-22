from datetime import date, datetime
from io import BytesIO
from unittest.mock import patch

import frappe
from frappe.utils import add_days, today
from openpyxl import load_workbook

from tracecore.core.errors import BlockedError
from tracecore.inventory import receipt_import
from tracecore.inventory.receipt_import import ImportFileError
from tracecore.tests import factories
from tracecore.tests.utils import IsolatedTestCase, bru

# Item/Batch/Stock Entry (ERPNext) kéo theo test record Company, Warehouse, Fiscal Year…, cần Setup Wizard;
# test này tự tạo dữ liệu qua factories nên bỏ phụ thuộc đó (site CI mới chưa có wizard).
test_ignore = [
	"Item",
	"Batch",
	"Stock Entry",
	"Warehouse",
	"Company",
	"Supplier",
	"Stock Entry Type",
	"Fiscal Year",
]

TOMORROW = add_days(today(), 1)
YESTERDAY = add_days(today(), -1)
FAR = add_days(today(), 400)


def cssd_item(code, group="Chemical"):
	"""Item vật tư như người dùng nhập trên Desk (luật FEAT-041 áp đủ)."""
	return factories.make_item(code, group, apply_rules=True)


def codes(issues):
	return sorted(issue.code for issue in issues)


class ImportCase(IsolatedTestCase):
	def setUp(self):
		super().setUp()
		self.masters = factories.ensure_stock_masters()
		self.warehouse = self.masters["warehouse"]
		self.supplier = factories.make_supplier("F41 NCC IMP")
		self.supplier_b = factories.make_supplier("F41 NCC IMP B")
		cssd_item("F41-I-1")
		cssd_item("F41-I-2")
		cssd_item("F41-I-POUCH", "Pouch")

	def row(self, item, lot, expiry=FAR, qty=10, rate=1000, supplier=None, doc="PGH-1", warehouse="default"):
		return [
			item,
			lot,
			expiry,
			qty,
			rate,
			self.supplier if supplier is None else supplier,
			doc,
			self.warehouse if warehouse == "default" else warehouse,
		]

	def workbook(self, rows, headers=None):
		return factories.make_receipt_workbook(rows, headers)

	def counts(self):
		return {
			"stock_entry": frappe.db.count("Stock Entry"),
			"batch": frappe.db.count("Batch"),
			"sle": frappe.db.count("Stock Ledger Entry"),
		}


class TestTemplateAndParse(ImportCase):
	def test_template_has_columns_and_guide_sheet(self):
		workbook = load_workbook(BytesIO(receipt_import.build_template()))
		self.assertEqual(workbook.sheetnames, ["Receipt", "Guide"])
		sheet = workbook["Receipt"]
		header = [cell.value for cell in sheet[1]]
		self.assertEqual(header, receipt_import.header_labels())
		self.assertEqual(len(header), len(receipt_import.TEMPLATE_COLUMNS))
		self.assertEqual(sheet["C2"].number_format, "dd/mm/yyyy")
		self.assertIsNone(sheet["A2"].value)  # không dữ liệu mẫu
		self.assertGreater(workbook["Guide"].max_row, 8)

	def test_template_roundtrip_parses_to_zero_rows(self):
		self.assertEqual(receipt_import.parse_workbook(receipt_import.build_template()), [])

	def test_parse_accepts_date_cell_and_dd_mm_yyyy_text(self):
		content = self.workbook(
			[
				self.row("F41-I-1", "L1", datetime(2027, 5, 31)),
				self.row("F41-I-1", "L2", "31/05/2027"),
				self.row("F41-I-1", "L3", "2027-05-31"),
				self.row("F41-I-1", "L4", date(2027, 5, 31)),
			]
		)
		rows = receipt_import.parse_workbook(content)
		self.assertEqual([r.expiry for r in rows], [date(2027, 5, 31)] * 4)
		self.assertTrue(all(not r.errors for r in rows))

	def test_parse_accepts_english_key_and_translated_headers(self):
		headers = list(receipt_import.TEMPLATE_COLUMNS)  # item_code, batch_id, …
		rows = receipt_import.parse_workbook(self.workbook([self.row("F41-I-1", "L1")], headers))
		self.assertEqual((rows[0].item_code, rows[0].batch_id, rows[0].qty), ("F41-I-1", "L1", 10))

	def test_parse_rejects_bad_dates_per_row(self):
		content = self.workbook(
			[
				self.row("F41-I-1", "L1", "31/02/2027"),
				self.row("F41-I-1", "L2", FAR),
				self.row("F41-I-1", "L3", "abc"),
				self.row("F41-I-1", "L4", 45000),  # số, không phải ngày
			]
		)
		_documents, issues = receipt_import.check_import(receipt_import.parse_workbook(content))
		bad = [(i.row, i.code, i.field) for i in issues if i.code == "DATE_INVALID"]
		self.assertEqual(
			bad,
			[
				(2, "DATE_INVALID", "expiry_date"),
				(4, "DATE_INVALID", "expiry_date"),
				(5, "DATE_INVALID", "expiry_date"),
			],
		)
		self.assertNotIn("EXPIRY_REQUIRED", codes(issues))  # lỗi ngày không kéo thêm lỗi thiếu hạn cùng dòng

	def test_row_numbers_match_excel_rows(self):
		content = self.workbook(
			[
				self.row("F41-I-1", "L1"),
				[None] * 8,  # dòng trống ở giữa vẫn chiếm số dòng
				self.row("F41-I-1", "L2"),
				[None] * 8,
			]
		)
		rows = receipt_import.parse_workbook(content)
		self.assertEqual([r.row for r in rows], [2, 4])

	def test_rejects_empty_file_wrong_extension_too_big_too_many_rows_missing_columns(self):
		valid = self.workbook([self.row("F41-I-1", "L1")])
		too_many = self.workbook([self.row("F41-I-1", f"L{n}") for n in range(receipt_import.MAX_ROWS + 1)])
		at_limit = self.workbook([self.row("F41-I-1", f"L{n}") for n in range(receipt_import.MAX_ROWS)])
		no_qty = [
			label for label in receipt_import.header_labels() if label != receipt_import.header_labels()[3]
		]
		cases = {
			"empty": b"",
			"not xlsx": b"item_code,batch_id\nA,B\n",
			"zip but not workbook": _zip_without_workbook(),
			"no header row": _empty_sheet_workbook(),
			"missing column": self.workbook([], no_qty),
			"too many rows": too_many,
		}
		for name, content in cases.items():
			with self.subTest(name), self.assertRaises(ImportFileError) as ctx:
				receipt_import.parse_workbook(content)
			self.assertTrue(ctx.exception.reason)
		with self.subTest("too big"), patch.object(receipt_import, "MAX_BYTES", len(valid) - 1):
			with self.assertRaises(ImportFileError):
				receipt_import.parse_workbook(valid)
		with self.subTest("exactly 1000 rows"):
			self.assertEqual(len(receipt_import.parse_workbook(at_limit)), receipt_import.MAX_ROWS)

	def test_header_only_file_has_no_data_rows(self):
		with self.assertRaises(ImportFileError):
			receipt_import.import_receipts(self.workbook([]), self.warehouse)


def _zip_without_workbook():
	import zipfile

	buffer = BytesIO()
	with zipfile.ZipFile(buffer, "w") as archive:
		archive.writestr("readme.txt", "x")
	return buffer.getvalue()


def _empty_sheet_workbook():
	from openpyxl import Workbook

	buffer = BytesIO()
	Workbook().save(buffer)
	return buffer.getvalue()


class TestCheckAndImport(ImportCase):
	@bru("BRU-INV-001")
	def test_invalid_rows_block_everything_and_report_all(self):
		before = self.counts()
		content = self.workbook(
			[
				self.row("F41-I-1", "OK-1"),  # dòng 2 đúng
				self.row("NO-SUCH-ITEM", "OK-2"),  # dòng 3 sai: Item không có
				self.row("F41-I-2", "OK-3"),  # dòng 4 đúng
				self.row("F41-I-1", "OK-4", YESTERDAY),  # dòng 5 sai: hết hạn
				self.row("F41-I-1", "", FAR),  # dòng 6 sai: thiếu lô
			]
		)
		with self.assertRaises(BlockedError) as ctx:
			receipt_import.import_receipts(content, self.warehouse)
		block = ctx.exception.block
		self.assertEqual(block.code, "IMPORT_ROWS_INVALID")
		self.assertEqual(block.http_status, 422)
		self.assertEqual(
			sorted((item["row"], item["code"]) for item in block.items),
			[(3, "ITEM_NOT_FOUND"), (5, "EXPIRY_PAST"), (6, "BATCH_REQUIRED")],
		)
		self.assertTrue(all(item["reason"] and item["field"] for item in block.items))
		self.assertEqual(self.counts(), before)  # không có Stock Entry/Batch/tồn mới nào
		self.assertFalse(frappe.db.exists("Batch", "OK-1"))

	def test_groups_rows_by_supplier_document_warehouse(self):
		other_warehouse = factories.make_warehouse("F41 Kho phu")
		content = self.workbook(
			[
				self.row("F41-I-1", "G-1", doc="D1"),
				self.row("F41-I-2", "G-2", doc="D1"),  # cùng nhóm với dòng trên
				self.row("F41-I-1", "G-3", doc="D2"),  # khác số chứng từ
				self.row("F41-I-1", "G-4", doc="D1", supplier=self.supplier_b),  # khác NCC
				self.row("F41-I-2", "G-5", doc="D1", warehouse=other_warehouse),  # khác kho
			]
		)
		result = receipt_import.import_receipts(content, self.warehouse)
		self.assertTrue(result["ok"])
		self.assertEqual(result["rows"], 5)
		self.assertEqual(len(result["stock_entries"]), 4)
		entries = [frappe.get_doc("Stock Entry", name) for name in result["stock_entries"]]
		self.assertTrue(all(e.docstatus == 1 and e.purpose == "Material Receipt" for e in entries))
		first = entries[0]
		self.assertEqual(
			(first.tc_supplier, first.tc_document_no, len(first.items)), (self.supplier, "D1", 2)
		)
		self.assertEqual(str(first.posting_date), today())
		self.assertEqual(
			frappe.db.get_value("Bin", {"item_code": "F41-I-2", "warehouse": other_warehouse}, "actual_qty"),
			10,
		)

	def test_same_lot_different_expiry_in_file_blocked(self):
		content = self.workbook(
			[
				self.row("F41-I-1", "DUP-1", FAR, doc="D1"),
				self.row("F41-I-1", "DUP-1", add_days(today(), 300), doc="D2"),  # nhóm khác, cùng file
			]
		)
		before = self.counts()
		with self.assertRaises(BlockedError) as ctx:
			receipt_import.import_receipts(content, self.warehouse)
		self.assertEqual(
			[(i["row"], i["code"]) for i in ctx.exception.block.items], [(3, "LOT_CONFLICT_IN_DOCUMENT")]
		)
		self.assertEqual(self.counts(), before)

	def test_existing_lot_with_other_expiry_or_item_blocked(self):
		factories.make_receipt(
			[{"item_code": "F41-I-1", "batch_id": "OLD-1", "expiry_date": FAR, "qty": 1}], self.supplier, "P0"
		)
		content = self.workbook(
			[
				self.row("F41-I-1", "OLD-1", add_days(today(), 100)),
				self.row("F41-I-2", "OLD-1", FAR),
			]
		)
		_documents, issues = receipt_import.check_import(receipt_import.parse_workbook(content))
		self.assertEqual(
			[(i.row, i.code) for i in issues], [(2, "BATCH_EXPIRY_MISMATCH"), (3, "BATCH_ITEM_MISMATCH")]
		)

	def test_default_warehouse_used_when_cell_empty_and_missing_both_blocked(self):
		content = self.workbook([self.row("F41-I-1", "W-1", warehouse=None)])
		result = receipt_import.import_receipts(content, self.warehouse)
		self.assertEqual(
			frappe.get_doc("Stock Entry", result["stock_entries"][0]).to_warehouse, self.warehouse
		)
		with self.assertRaises(BlockedError) as ctx:
			receipt_import.import_receipts(self.workbook([self.row("F41-I-1", "W-2", warehouse=None)]), None)
		self.assertEqual(
			[(i["row"], i["code"]) for i in ctx.exception.block.items], [(2, "WAREHOUSE_REQUIRED")]
		)

	def test_supplier_document_and_warehouse_must_exist(self):
		group = factories.make_warehouse("F41 Nhom kho", is_group=True)
		content = self.workbook(
			[
				self.row("F41-I-1", "S-1", supplier="Khong co NCC"),
				self.row("F41-I-1", "S-2", warehouse="Khong co kho - TCT"),
				self.row("F41-I-1", "S-3", warehouse=group),
				self.row("F41-I-1", "S-4", doc=""),
				self.row("F41-I-1", "S-5", supplier=""),
			]
		)
		_documents, issues = receipt_import.check_import(receipt_import.parse_workbook(content))
		self.assertEqual(
			[(i.row, i.code) for i in issues],
			[
				(2, "SUPPLIER_NOT_FOUND"),
				(3, "WAREHOUSE_NOT_FOUND"),
				(4, "WAREHOUSE_NOT_FOUND"),
				(5, "DOCUMENT_REQUIRED"),
				(6, "SUPPLIER_REQUIRED"),
			],
		)

	def test_expired_and_missing_expiry_rows_blocked_like_desk(self):
		content = self.workbook(
			[
				self.row("F41-I-1", "E-1", YESTERDAY),
				self.row("F41-I-1", "E-2", None),
				self.row("F41-I-1", "E-3", date.fromisoformat(today())),  # hôm nay vẫn nhận
				self.row("F41-I-POUCH", "E-4", FAR),  # không theo dõi hạn mà có hạn
				self.row("F41-I-1", "E-5", FAR, qty=0),
				self.row("F41-I-1", "E-6", FAR, qty="abc"),
				self.row("F41-I-1", "E-7", FAR, rate=-1),
				self.row("F41-I-1", "E-8", FAR, rate="x"),
			]
		)
		_documents, issues = receipt_import.check_import(receipt_import.parse_workbook(content))
		self.assertEqual(
			[(i.row, i.code) for i in issues],
			[
				(2, "EXPIRY_PAST"),
				(3, "EXPIRY_REQUIRED"),
				(5, "EXPIRY_NOT_TRACKED"),
				(6, "QTY_INVALID"),
				(7, "QTY_INVALID"),
				(8, "RATE_NEGATIVE"),
				(9, "RATE_INVALID"),
			],
		)

	def test_failure_during_submit_rolls_back_all_documents(self):
		content = self.workbook(
			[self.row("F41-I-1", "RB-1", doc="D1"), self.row("F41-I-2", "RB-2", doc="D2")]
		)
		before = self.counts()
		real = receipt_import._post_entry
		calls = []

		def flaky(entry):
			calls.append(entry)
			if len(calls) == 2:
				raise RuntimeError("boom")
			return real(entry)

		with patch.object(receipt_import, "_post_entry", flaky), self.assertRaises(RuntimeError):
			receipt_import.import_receipts(content, self.warehouse)
		self.assertEqual(len(calls), 2)
		self.assertEqual(self.counts(), before)
		self.assertFalse(frappe.db.exists("Batch", "RB-1"))
		self.assertIsNone(
			frappe.db.get_value("Bin", {"item_code": "F41-I-1", "warehouse": self.warehouse}, "actual_qty")
		)

	def test_dry_run_writes_nothing(self):
		before = self.counts()
		content = self.workbook([self.row("F41-I-1", "DR-1"), self.row("F41-I-2", "DR-2", doc="D2", rate=0)])
		result = receipt_import.import_receipts(content, self.warehouse, dry_run=True)
		self.assertEqual((result["ok"], result["rows"], result["documents"]), (True, 2, 2))
		self.assertEqual([w["row"] for w in result["warnings"]], [3])  # đơn giá 0: cảnh báo, không chặn
		self.assertEqual(self.counts(), before)

	def test_company_not_configured_blocks(self):
		content = self.workbook([self.row("F41-I-1", "C-1")])
		with patch("erpnext.get_default_company", return_value=None), self.assertRaises(BlockedError) as ctx:
			receipt_import.import_receipts(content, self.warehouse)
		self.assertEqual(
			(ctx.exception.block.code, ctx.exception.block.http_status), ("COMPANY_NOT_CONFIGURED", 409)
		)
