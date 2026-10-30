"""Nhập kho hàng loạt từ file Excel mẫu (FEAT-041 Task 4, FR-M6-01.2, FR-M9-11, A-07 V12).

Một file ``.xlsx`` (≤ 2 MB, ≤ 1000 dòng dữ liệu, sheet đầu, dòng 1 là tiêu đề). **Tất cả hoặc không gì:** còn một
dòng lỗi thì không ghi gì và báo MỌI dòng lỗi theo số dòng nhìn thấy trong Excel (tiêu đề = dòng 1). Luật từng dòng
là ``receipt.check_lines`` / ``receipt.check_header`` (cùng hàm với form Desk), file này chỉ thêm phần đọc file, kiểm
Nhà cung cấp / Kho / ngày / số, gom dòng theo (NCC, số chứng từ, kho) thành mỗi nhóm một ``Stock Entry`` đã ghi sổ.
"""

import math
import zipfile
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import cint
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from tracecore.core.errors import Block, BlockedError
from tracecore.inventory.receipt import (
	LineIssue,
	ReceiptInvalid,
	ReceiptLine,
	check_header,
	check_lines,
	today_vn,
	zero_rate_rows,
)

TEMPLATE_COLUMNS = (
	"item_code",
	"batch_id",
	"expiry_date",
	"qty",
	"rate",
	"supplier",
	"document_no",
	"warehouse",
)
REQUIRED_COLUMNS = TEMPLATE_COLUMNS[:-1]  # kho có thể lấy từ kho mặc định chọn trong hộp thoại
MAX_ROWS = 1000
MAX_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024  # chặn zip bomb: file nén nhỏ nhưng giải nén rất lớn
MAX_SCAN_ROWS = 10_000  # số dòng vật lý tối đa được quét (kể cả dòng trống có định dạng)
DATE_FORMAT = "dd/mm/yyyy"
SAVEPOINT = "tc_receipt_import"

# Tiêu đề tiếng Anh (luôn được nhận) cạnh bản dịch theo ngôn ngữ người dùng (``header_labels``).
ENGLISH_LABELS = {
	"item_code": "Item code",
	"batch_id": "Lot number",
	"expiry_date": "Expiry date",
	"qty": "Received quantity",
	"rate": "Unit price",
	"supplier": "Supplier name",
	"document_no": "Supplier document number",
	"warehouse": "Receiving warehouse",
}


class ImportFileError(Exception):
	"""File không đọc được hoặc vượt giới hạn (``IMPORT_FILE_INVALID``); ``reason`` đã qua ``_()``."""

	def __init__(self, reason: str):
		super().__init__(reason)
		self.reason = reason


@dataclass
class ImportRow:
	row: int  # số dòng Excel (tiêu đề = 1)
	item_code: str
	batch_id: str | None
	expiry: date | None
	qty: float | None
	rate: float
	supplier: str
	document_no: str
	warehouse: str | None
	errors: list[tuple[str, str, str]] = field(default_factory=list)  # (mã lỗi, cột, lý do)


@dataclass
class ImportDocument:
	supplier: str
	document_no: str
	warehouse: str
	lines: list[ReceiptLine]


def _labels() -> dict[str, str]:
	return {
		"item_code": _("Item code"),
		"batch_id": _("Lot number"),
		"expiry_date": _("Expiry date"),
		"qty": _("Received quantity"),
		"rate": _("Unit price"),
		"supplier": _("Supplier name"),
		"document_no": _("Supplier document number"),
		"warehouse": _("Receiving warehouse"),
	}


def header_labels() -> list[str]:
	"""Tiêu đề cột của file mẫu theo ngôn ngữ hiện tại, đúng thứ tự ``TEMPLATE_COLUMNS``."""
	labels = _labels()
	return [labels[code] for code in TEMPLATE_COLUMNS]


def _norm(value) -> str:
	return " ".join(str(value).split()).casefold() if value is not None else ""


def _header_aliases() -> dict[str, str]:
	"""Tiêu đề chuẩn hóa -> mã cột; nhận mã cột, tiêu đề tiếng Anh và bản dịch hiện tại."""
	aliases: dict[str, str] = {}
	labels = _labels()
	for code in TEMPLATE_COLUMNS:
		for text in (code, ENGLISH_LABELS[code], labels[code]):
			aliases.setdefault(_norm(text), code)
	return aliases


def _guide_rows() -> list[tuple[str, str, str]]:
	labels = _labels()
	required, optional = _("Required"), _("Optional")
	return [
		(labels["item_code"], required, _("Code of a CSSD consumable item that already exists in ERPNext.")),
		(
			labels["batch_id"],
			required,
			_("Lot number printed on the package, up to 140 characters. One lot has one expiry date."),
		),
		(
			labels["expiry_date"],
			_("Required for chemicals and indicators"),
			_(
				"Date cell or text dd/mm/yyyy. Required for chemical indicators, biological indicators and chemicals; must not be earlier than today; leave empty for items that do not track expiry."
			),
		),
		(labels["qty"], required, _("Number greater than 0.")),
		(labels["rate"], optional, _("Price per unit in VND. Empty or 0 is received with a warning.")),
		(labels["supplier"], required, _("Name of an existing supplier.")),
		(
			labels["document_no"],
			required,
			_(
				"Delivery note or invoice number. Rows with the same supplier, document number and warehouse form one receipt."
			),
		),
		(
			labels["warehouse"],
			optional,
			_("Name of a receiving warehouse. Empty uses the warehouse chosen when importing."),
		),
	]


def build_template() -> bytes:
	"""File mẫu ``.xlsx`` hai sheet: ``Receipt`` (tiêu đề, không dữ liệu mẫu) và ``Guide`` (mô tả cột, quy tắc)."""
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "Receipt"
	sheet.append(header_labels())
	bold = Font(bold=True)
	for cell in sheet[1]:
		cell.font = bold
	expiry_col = TEMPLATE_COLUMNS.index("expiry_date") + 1
	for row in range(2, MAX_ROWS + 2):
		sheet.cell(row=row, column=expiry_col).number_format = DATE_FORMAT
	for index in range(len(TEMPLATE_COLUMNS)):
		sheet.column_dimensions[chr(ord("A") + index)].width = 28
	sheet.freeze_panes = "A2"

	guide = workbook.create_sheet("Guide")
	guide.append([_("Excel column"), _("Requirement"), _("Description")])
	for cell in guide[1]:
		cell.font = bold
	for values in _guide_rows():
		guide.append(list(values))
	guide.append([])
	for rule in (
		_("Fill in the Receipt sheet from row 2; row 1 is the header and must not be changed."),
		_("At most {0} rows and {1} MB per file.").format(MAX_ROWS, MAX_BYTES // (1024 * 1024)),
		_("If any row is wrong nothing is saved; every wrong row is listed so you can fix them all at once."),
		_("Posting date is today; receipts are posted as submitted Stock Entries."),
	):
		guide.append([rule])
	guide.column_dimensions["A"].width = 30
	guide.column_dimensions["B"].width = 34
	guide.column_dimensions["C"].width = 90
	buffer = BytesIO()
	workbook.save(buffer)
	return buffer.getvalue()


# ---------------------------------------------------------------------------------------------------------------------
# Đọc file
# ---------------------------------------------------------------------------------------------------------------------


def _file_error(reason: str) -> ImportFileError:
	return ImportFileError(reason + " " + _("Download the template, fill it in and upload it again."))


def _check_container(content: bytes) -> None:
	if not content:
		raise _file_error(_("The file is empty."))
	if len(content) > MAX_BYTES:
		raise _file_error(_("The file is larger than {0} MB.").format(MAX_BYTES // (1024 * 1024)))
	try:
		with zipfile.ZipFile(BytesIO(content)) as archive:
			total = sum(info.file_size for info in archive.infolist())
			names = set(archive.namelist())
	except zipfile.BadZipFile:
		raise _file_error(_("The file is not an Excel .xlsx workbook.")) from None
	if "xl/workbook.xml" not in names:
		raise _file_error(_("The file is not an Excel .xlsx workbook."))
	if total > MAX_UNCOMPRESSED_BYTES:
		raise _file_error(_("The workbook is too large to read."))


def _text(value) -> str:
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		return str(int(value))
	return str(value).strip()


def _shown(value) -> str:
	return _text(value)[:40]


def _parse_date(value) -> tuple[date | None, bool]:
	"""(ngày, hợp lệ). Ô ngày, hoặc chuỗi ``dd/mm/yyyy`` / ``yyyy-mm-dd``; còn lại không hợp lệ."""
	if value is None or (isinstance(value, str) and not value.strip()):
		return None, True
	if isinstance(value, datetime):
		return value.date(), True
	if isinstance(value, date):
		return value, True
	if isinstance(value, str):
		text = value.strip()
		for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
			try:
				return datetime.strptime(text, pattern).date(), True
			except ValueError:
				continue
	return None, False


def _parse_number(value) -> tuple[float | None, bool]:
	"""(số, hợp lệ); ô trống cho (None, True)."""
	if value is None or (isinstance(value, str) and not value.strip()):
		return None, True
	if isinstance(value, bool):
		return None, False
	if isinstance(value, int | float):
		number = float(value)
	elif isinstance(value, str):
		try:
			number = float(value.strip().replace(",", "."))
		except ValueError:
			return None, False
	else:
		return None, False
	return (number, True) if math.isfinite(number) else (None, False)


def _build_row(number: int, cells: dict[str, object]) -> ImportRow:
	errors: list[tuple[str, str, str]] = []
	expiry, ok = _parse_date(cells.get("expiry_date"))
	if not ok:
		errors.append(
			(
				"DATE_INVALID",
				"expiry_date",
				_(
					'Expiry date "{0}" is not a valid date. Enter it as dd/mm/yyyy, or use a date cell.'
				).format(_shown(cells.get("expiry_date"))),
			)
		)
	qty, _ok = _parse_number(cells.get("qty"))  # sai/rỗng -> None; check_lines báo QTY_INVALID
	rate, ok = _parse_number(cells.get("rate"))
	if not ok:
		errors.append(
			(
				"RATE_INVALID",
				"rate",
				_(
					'Unit price "{0}" is not a number. Enter a number such as 125000, or leave it empty.'
				).format(_shown(cells.get("rate"))),
			)
		)
	return ImportRow(
		row=number,
		item_code=_text(cells.get("item_code")),
		batch_id=_text(cells.get("batch_id")) or None,
		expiry=expiry,
		qty=qty,
		rate=rate or 0.0,
		supplier=_text(cells.get("supplier")),
		document_no=_text(cells.get("document_no")),
		warehouse=_text(cells.get("warehouse")) or None,
		errors=errors,
	)


def parse_workbook(content: bytes) -> list[ImportRow]:
	"""Đọc sheet đầu của ``.xlsx``; trả các dòng dữ liệu (bỏ dòng trống hoàn toàn) với số dòng Excel.

	Ném ``ImportFileError`` khi file rỗng / không phải .xlsx / quá cỡ / quá ``MAX_ROWS`` dòng / thiếu cột bắt buộc.
	Ô công thức lấy giá trị đã lưu (``data_only``); công thức chưa tính coi như ô trống.
	"""
	_check_container(content)
	try:
		workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
	except Exception:
		raise _file_error(_("The file is not an Excel .xlsx workbook.")) from None
	try:
		return _read_sheet(workbook)
	except ImportFileError:
		raise
	except Exception:
		raise _file_error(_("The workbook could not be read.")) from None
	finally:
		workbook.close()


def _read_sheet(workbook) -> list[ImportRow]:
	if not workbook.worksheets:
		raise _file_error(_("The workbook has no sheet."))
	iterator = workbook.worksheets[0].iter_rows(values_only=True)
	header = next(iterator, None)
	if not header or all(_text(cell) == "" for cell in header):
		raise _file_error(_("Row 1 of the first sheet must be the header row."))
	aliases = _header_aliases()
	columns: dict[str, int] = {}
	for index, cell in enumerate(header):
		code = aliases.get(_norm(cell))
		if code:
			columns.setdefault(code, index)
	missing = [code for code in REQUIRED_COLUMNS if code not in columns]
	if missing:
		labels = _labels()
		raise _file_error(_("Missing column(s): {0}.").format(", ".join(labels[code] for code in missing)))

	rows: list[ImportRow] = []
	for number, values in enumerate(iterator, start=2):
		if number > MAX_SCAN_ROWS:
			raise _file_error(_("The sheet has too many rows to read."))
		if all(_text(cell) == "" for cell in values):
			continue
		if len(rows) >= MAX_ROWS:
			raise _file_error(_("The file has more than {0} data rows.").format(MAX_ROWS))
		cells = {code: (values[index] if index < len(values) else None) for code, index in columns.items()}
		rows.append(_build_row(number, cells))
	return rows


# ---------------------------------------------------------------------------------------------------------------------
# Kiểm
# ---------------------------------------------------------------------------------------------------------------------


def _default_company() -> str | None:
	import erpnext

	return erpnext.get_default_company()


def _lookup(cache: dict, key: str, loader):
	if key not in cache:
		cache[key] = loader(key)
	return cache[key]


def check_import(
	rows: list[ImportRow], default_warehouse: str | None = None
) -> tuple[list[ImportDocument], list[LineIssue]]:
	"""Kiểm mọi dòng (không ghi DB); trả (các nhóm chứng từ, TẤT CẢ lỗi sắp theo số dòng).

	``check_lines`` chạy một lần trên cả file nên lô trùng giữa hai nhóm khác hạn vẫn bị bắt.
	"""
	company = _default_company()
	issues: list[LineIssue] = []
	lines: list[ReceiptLine] = []
	date_failed: set[int] = set()
	groups: dict[tuple[str, str, str], ImportDocument] = {}
	suppliers: dict = {}
	warehouses: dict = {}

	for item in rows:
		for code, column, reason in item.errors:
			issues.append(LineIssue(item.row, code, column, reason))
			if code == "DATE_INVALID":
				date_failed.add(item.row)
		warehouse = item.warehouse or default_warehouse
		issues += [
			replace(issue, row=item.row) for issue in check_header(item.supplier, item.document_no, warehouse)
		]
		if item.supplier and not _lookup(suppliers, item.supplier, _supplier_usable):
			issues.append(
				LineIssue(
					item.row,
					"SUPPLIER_NOT_FOUND",
					"supplier",
					_(
						"Supplier {0} does not exist or is disabled. Check the supplier name or create the supplier first."
					).format(item.supplier),
				)
			)
		if warehouse:
			problem = _lookup(warehouses, warehouse, lambda name: _warehouse_problem(name, company))
			if problem:
				issues.append(LineIssue(item.row, "WAREHOUSE_NOT_FOUND", "warehouse", problem))
		line = ReceiptLine(
			row=item.row,
			item_code=item.item_code,
			batch_id=item.batch_id,
			expiry=item.expiry,
			qty=item.qty or 0.0,
			rate=item.rate,
		)
		lines.append(line)
		if item.supplier and item.document_no and warehouse:
			key = (item.supplier, item.document_no, warehouse)
			groups.setdefault(key, ImportDocument(*key, lines=[])).lines.append(line)

	issues += [
		issue
		for issue in check_lines(lines, today_vn())
		if not (issue.code == "EXPIRY_REQUIRED" and issue.row in date_failed)
	]
	issues.sort(key=lambda issue: issue.row)  # sắp ổn định: cùng dòng giữ thứ tự phát hiện
	return list(groups.values()), issues


def _supplier_usable(name: str) -> bool:
	disabled = frappe.db.get_value("Supplier", name, "disabled")
	return disabled is not None and not cint(disabled)


def _warehouse_problem(name: str, company: str | None) -> str | None:
	values = frappe.db.get_value("Warehouse", name, ["is_group", "company", "disabled"], as_dict=True)
	if not values or cint(values.disabled):
		return _("Warehouse {0} does not exist or is disabled. Choose an existing warehouse.").format(name)
	if cint(values.is_group):
		return _("Warehouse {0} is a group. Choose a warehouse that can hold stock.").format(name)
	if company and values.company != company:
		return _("Warehouse {0} belongs to another company. Choose a warehouse of {1}.").format(name, company)
	return None


# ---------------------------------------------------------------------------------------------------------------------
# Khối lỗi (Block) cho tầng RPC
# ---------------------------------------------------------------------------------------------------------------------


def file_block(reason: str) -> Block:
	return Block(
		"IMPORT_FILE_INVALID",
		_("The Excel file cannot be used"),
		reason,
		_("Fix the file, then check it again."),
	)


def company_block() -> Block:
	return Block(
		"COMPANY_NOT_CONFIGURED",
		_("Cannot receive these items"),
		_("No company is configured for this site. Ask the administrator to finish the ERPNext setup."),
		_("Ask the administrator to finish the ERPNext setup, then import the file again."),
	)


def rows_block(issues: list[LineIssue]) -> Block:
	rows = sorted({issue.row for issue in issues})
	return Block(
		"IMPORT_ROWS_INVALID",
		_("The Excel file has errors; nothing was saved"),
		_("{0} problem(s) in {1} row(s) stop this file from being imported.").format(len(issues), len(rows)),
		_("Fix the rows listed, then check the file again."),
		items=[
			{"row": issue.row, "code": issue.code, "field": issue.field, "reason": issue.reason}
			for issue in issues
		],
	)


# ---------------------------------------------------------------------------------------------------------------------
# Ghi sổ
# ---------------------------------------------------------------------------------------------------------------------


def _build_entry(company: str, document: ImportDocument):
	items = []
	for line in document.lines:
		stock_uom = frappe.db.get_value("Item", line.item_code, "stock_uom")
		items.append(
			{
				"item_code": line.item_code,
				"qty": line.qty,
				"basic_rate": line.rate,
				"uom": stock_uom,
				"stock_uom": stock_uom,
				"conversion_factor": 1,
				"t_warehouse": document.warehouse,
				"tc_batch_id": line.batch_id,
				"tc_expiry_date": line.expiry,
			}
		)
	return frappe.get_doc(
		{
			"doctype": "Stock Entry",
			"stock_entry_type": "Material Receipt",
			"purpose": "Material Receipt",
			"company": company,
			"posting_date": today_vn(),
			"to_warehouse": document.warehouse,
			"tc_supplier": document.supplier,
			"tc_document_no": document.document_no,
			"items": items,
		}
	)


def _post_entry(entry) -> str:
	"""Lập và ghi sổ một phiếu; kiểm quyền của người gọi (không ``ignore_permissions``)."""
	entry.insert()
	entry.submit()
	return entry.name


def import_receipts(content: bytes, default_warehouse: str | None = None, dry_run: bool = False) -> dict:
	"""Kiểm (``dry_run``) hoặc kiểm rồi ghi sổ toàn bộ file; tất cả hoặc không gì.

	Thành công: ``{"ok": True, "rows", "documents"|"stock_entries", "warnings"}``. Không ghi: ``ImportFileError``
	(file), ``BlockedError`` mang ``IMPORT_ROWS_INVALID`` (kèm mọi dòng lỗi) hoặc ``COMPANY_NOT_CONFIGURED``.
	"""
	company = _default_company()
	if not company:
		raise BlockedError(company_block())
	rows = parse_workbook(content)
	if not rows:
		raise _file_error(_("The file has no data rows."))
	documents, issues = check_import(rows, default_warehouse)
	if issues:
		raise BlockedError(rows_block(issues))
	warnings = [
		{
			"row": number,
			"code": "RATE_ZERO",
			"reason": _("Unit price is 0; the stock is received with a zero valuation rate."),
		}
		for number in sorted(zero_rate_rows([line for document in documents for line in document.lines]))
	]
	if dry_run:
		return {"ok": True, "rows": len(rows), "documents": len(documents), "warnings": warnings}

	logged = len(frappe.local.message_log)
	frappe.db.savepoint(SAVEPOINT)
	names: list[str] = []
	try:
		for document in documents:
			excel_rows = [line.row for line in document.lines]
			try:
				names.append(_post_entry(_build_entry(company, document)))
			except (
				ReceiptInvalid
			) as error:  # luật đổi giữa lúc kiểm và ghi: ánh xạ số dòng phiếu về số dòng Excel
				raise BlockedError(rows_block(_remap(error.issues, excel_rows))) from error
	except BaseException:
		frappe.db.rollback(save_point=SAVEPOINT)
		del frappe.local.message_log[logged:]  # lỗi đi qua body RPC, không bật hộp thoại Desk
		raise
	del frappe.local.message_log[logged:]
	return {
		"ok": True,
		"rows": len(rows),
		"stock_entries": names,
		"warnings": warnings,
	}


def _remap(issues: list[LineIssue], excel_rows: list[int]) -> list[LineIssue]:
	return [
		replace(issue, row=excel_rows[issue.row - 1]) if 0 < issue.row <= len(excel_rows) else issue
		for issue in issues
	]


__all__ = [
	"MAX_BYTES",
	"MAX_ROWS",
	"TEMPLATE_COLUMNS",
	"ImportFileError",
	"ImportRow",
	"build_template",
	"check_import",
	"header_labels",
	"import_receipts",
	"parse_workbook",
]
