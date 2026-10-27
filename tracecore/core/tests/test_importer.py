"""Khung import danh mục từ Excel (FEAT-011 Task 2; FR-M9-11, EI-SW-007; A-41 Q2, Q3, D1-D3, D5)."""

from datetime import date
from io import BytesIO
from unittest.mock import patch

import frappe
from openpyxl import Workbook, load_workbook

from tracecore.core import importer
from tracecore.core.importer import Column, ImportFileError, ImportKind, Issue, RowError
from tracecore.tests.utils import IsolatedTestCase, bru

KEY = "test_note"


def _exists(code):
	return bool(frappe.db.exists("Note", {"title": code}))


def _validate(unit):
	issues = []
	for row in unit:
		if row.values.get("count") is not None and row.values["count"] < 0:
			issues.append(Issue(row.row, "COUNT_INVALID", "count", "count must be >= 0"))
	return issues


def _write(unit):
	head = unit[0].values
	if head["code"] == "BOOM":
		raise RowError([Issue(unit[0].row, "WRITE_FAILED", None, "boom")])
	if head["code"] == "CRASH":
		frappe.throw("crash")
	lines = ", ".join(str(row.values.get("tags") or []) for row in unit)
	doc = frappe.get_doc({"doctype": "Note", "title": head["code"], "content": lines, "public": 1})
	doc.insert(ignore_permissions=True)
	return doc.name


COLUMNS = [
	Column("code", "Code", required=True),
	Column("count", "Count", kind="int"),
	Column("on", "Date", kind="date"),
	Column("flag", "Flag", kind="check"),
	Column("color", "Color", kind="select", options=("Red", "Blue")),
	Column("tags", "Tags", kind="list", options=("A", "B")),
]


def fake_kind(group=False):
	return ImportKind(
		key=KEY,
		label="Test notes",
		doctype="Note",
		columns=COLUMNS,
		code_field="code",
		exists=_exists,
		validate=_validate,
		write=_write,
		group_by="code" if group else None,
		group_fields=("color",) if group else (),
		forbidden_columns=("Set ID",),
	)


def workbook(header, rows):
	book = Workbook()
	sheet = book.active
	sheet.append(header)
	for row in rows:
		sheet.append(row)
	buffer = BytesIO()
	book.save(buffer)
	return buffer.getvalue()


class ImporterCase(IsolatedTestCase):
	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")
		self.kind = fake_kind()
		patcher = patch.object(importer, "kinds", side_effect=lambda: {KEY: self.kind})
		patcher.start()
		self.addCleanup(patcher.stop)


class TestTemplateAndParse(ImporterCase):
	def test_template_has_required_columns_and_help_sheet(self):
		book = load_workbook(BytesIO(importer.build_template(KEY)))
		data, guide = book.worksheets[0], book.worksheets[1]
		header = [cell.value for cell in data[1]]
		self.assertEqual(header, ["Code *", "Count", "Date", "Flag", "Color", "Tags"])
		self.assertTrue(data["A1"].font.bold)
		text = " ".join(str(cell.value) for row in guide.iter_rows() for cell in row if cell.value)
		self.assertIn("Red", text)  # giá trị hợp lệ của cột chọn
		self.assertIn("dd/mm/yyyy", text)

	def test_parse_by_header_name_any_order(self):
		content = workbook(["Color", "code", "Count *"], [["red", "n-1", 3]])
		rows, issues = importer.parse(self.kind, content)
		self.assertEqual(issues, [])
		self.assertEqual(rows[0].row, 2)
		self.assertEqual(rows[0].values["code"], "n-1")
		self.assertEqual(rows[0].values["color"], "Red")  # chuẩn hóa theo giá trị hợp lệ
		self.assertEqual(rows[0].values["count"], 3)

	def test_missing_required_column_rejects_file(self):
		with self.assertRaises(ImportFileError) as caught:
			importer.parse(self.kind, workbook(["Count"], [[1]]))
		self.assertIn("Code", caught.exception.reason)

	@bru("BRU-ID-006")
	def test_system_code_column_rejects_file(self):
		with self.assertRaises(ImportFileError) as caught:
			importer.parse(self.kind, workbook(["Code", "Set ID"], [["N1", "9520000ABC000001"]]))
		self.assertIn("Set ID", caught.exception.reason)

	def test_cell_normalization(self):
		content = workbook(
			["Code", "Count", "Date", "Flag", "Tags"],
			[
				[1234.0, "7", date(2026, 10, 27), "có", "a; b"],
				["  N 2 ", 2.0, "27/10/2026", "x", ""],
				["N3", "abc", "31/02/2026", "maybe", "C"],
			],
		)
		rows, issues = importer.parse(self.kind, content)
		self.assertEqual(rows[0].values["code"], "1234")  # số trong ô mã không thành "1234.0"
		self.assertEqual(rows[0].values["count"], 7)
		self.assertEqual(rows[0].values["on"], date(2026, 10, 27))
		self.assertIs(rows[0].values["flag"], True)
		self.assertEqual(rows[0].values["tags"], ["A", "B"])
		self.assertEqual(rows[1].values["code"], "N 2")
		self.assertEqual(rows[1].values["on"], date(2026, 10, 27))
		self.assertEqual(rows[1].values["tags"], [])
		self.assertEqual(
			sorted((issue.row, issue.field) for issue in issues),
			[(4, "count"), (4, "flag"), (4, "on"), (4, "tags")],
		)

	def test_not_a_workbook_rejected(self):
		with self.assertRaises(ImportFileError):
			importer.parse(self.kind, b"not an excel file")


class TestCheckAndRun(ImporterCase):
	def test_check_writes_nothing(self):
		before = frappe.db.count("Note")
		result = importer.check(KEY, workbook(["Code", "Count"], [["N1", 1], ["N2", -1]]))
		self.assertEqual(frappe.db.count("Note"), before)
		self.assertEqual((result["total"], result["valid"]), (2, 1))
		self.assertEqual([(i["row"], i["code"]) for i in result["issues"]], [(3, "COUNT_INVALID")])

	def test_existing_codes_are_skipped(self):
		frappe.get_doc({"doctype": "Note", "title": "OLD", "public": 1}).insert(ignore_permissions=True)
		content = workbook(["Code"], [["OLD"], ["NEW"]])
		self.assertEqual(importer.check(KEY, content)["existing"], ["OLD"])
		batch = frappe.get_doc("Import Batch", importer.run(KEY, content, "x.xlsx"))
		self.assertEqual((batch.created, batch.skipped, batch.failed), (1, 1, 0))
		self.assertEqual(frappe.db.count("Note", {"title": "OLD"}), 1)

	def test_failure_in_one_unit_keeps_others(self):
		content = workbook(["Code", "Count"], [["A1", 1], ["BOOM", 1], ["CRASH", 1], ["A2", -5], ["A3", 2]])
		batch = frappe.get_doc("Import Batch", importer.run(KEY, content, "x.xlsx"))
		self.assertEqual((batch.status, batch.created, batch.failed), ("Done", 2, 3))
		self.assertTrue(frappe.db.exists("Note", {"title": "A1"}))
		self.assertTrue(frappe.db.exists("Note", {"title": "A3"}))
		self.assertFalse(frappe.db.exists("Note", {"title": "CRASH"}))
		errors = {(row.row_no, row.result) for row in batch.rows if row.result == "Error"}
		self.assertEqual(errors, {(3, "Error"), (4, "Error"), (5, "Error")})

	def test_batch_records_row_results(self):
		batch = frappe.get_doc("Import Batch", importer.run(KEY, workbook(["Code"], [["R1"]]), "notes.xlsx"))
		self.assertEqual(batch.kind, KEY)
		row = batch.rows[0]
		self.assertEqual((row.row_no, row.code, row.result, row.ref_doctype), (2, "R1", "Created", "Note"))
		self.assertTrue(frappe.db.exists("Note", row.ref_name))
		self.assertTrue(batch.finished_at)

	def test_large_file_is_queued(self):
		rows = [[f"L{i}"] for i in range(importer.BACKGROUND_ROWS + 1)]
		with patch("frappe.enqueue") as enqueue:
			name = importer.run(KEY, workbook(["Code"], rows), "big.xlsx")
		self.assertEqual(frappe.db.get_value("Import Batch", name, "status"), "Queued")
		enqueue.assert_called_once()
		self.assertEqual(enqueue.call_args.kwargs["queue"], "long")
		self.assertEqual(enqueue.call_args.kwargs["batch"], name)
		importer.execute_batch(name, enqueue.call_args.kwargs["content"])  # như job nền nhận từ hàng đợi
		self.assertEqual(frappe.db.get_value("Import Batch", name, "created"), importer.BACKGROUND_ROWS + 1)


class TestGroups(ImporterCase):
	def setUp(self):
		super().setUp()
		self.kind = fake_kind(group=True)

	def test_group_rows_become_one_unit(self):
		content = workbook(
			["Code", "Color", "Tags"], [["G1", "Red", "A"], ["G1", "", "B"], ["G2", "Blue", "A"]]
		)
		batch = frappe.get_doc("Import Batch", importer.run(KEY, content, "g.xlsx"))
		self.assertEqual(batch.created, 2)
		self.assertEqual(frappe.db.count("Note", {"title": ["in", ["G1", "G2"]]}), 2)

	def test_group_rows_header_conflict(self):
		content = workbook(["Code", "Color"], [["G1", "Red"], ["G1", "Blue"]])
		result = importer.check(KEY, content)
		self.assertEqual(result["valid"], 0)
		self.assertEqual([(i["row"], i["code"]) for i in result["issues"]], [(3, "GROUP_CONFLICT")])
